from __future__ import annotations

import unittest
import ast
from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation
from pptx.util import Inches

from ppt_system.export.editable_charts import add_editable_chart
from ppt_system.export.native_shapes import add_native_shapes
from ppt_system.export.text_script_runtime import build_project_script_source


class NativeShapesTests(unittest.TestCase):
    def test_native_shape_and_chart_are_editable_objects(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        page = {
            "native_blueprint": [
                {"kind": "rect", "left": 100, "top": 100, "width": 500, "height": 260, "fill": "DDEBFF"}
            ]
        }
        scale = lambda value: Inches(value / 150)
        add_native_shapes(slide, page, scale, scale, scale, scale)
        add_editable_chart(
            slide,
            {"type": "bar", "categories": ["Q1", "Q2"], "series": [{"name": "收入", "values": [10, 20]}]},
            Inches(5), Inches(1), Inches(5), Inches(3),
        )
        self.assertEqual(len(slide.shapes), 2)
        self.assertFalse(slide.shapes[0].shape_type is None)
        self.assertTrue(slide.shapes[1].has_chart)

    def test_generated_script_uses_python_boolean_literals(self):
        project = {
            "image_width": 2000,
            "image_height": 1125,
            "pages": [{"page_no": 1, "title": "测试", "texts": [], "locked": True,
                       "chart_data": {"type": "bar", "categories": ["A"],
                                      "series": [{"name": "值", "values": [1]}],
                                      "show_legend": False}}],
        }
        with TemporaryDirectory() as directory:
            source = build_project_script_source(
                project, Path(directory), Path(directory) / "result.pptx",
                [{"page_no": 1, "script": "pass"}], layer_mode="hybrid_editable",
            )
        ast.parse(source)
        self.assertIn("'locked': True", source)
        self.assertIn("'show_legend': False", source)
        self.assertIn("page_texts = PAGE_TEXTS[1]", source)
        self.assertNotIn('page_texts = PAGE_TEXTS["1"]', source)


if __name__ == "__main__":
    unittest.main()
