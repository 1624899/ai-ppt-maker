from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from ppt_system.export.delivery_options import EDITABLE_PPT_FILENAMES
from ppt_system.export.export_layer_mode import OVERLAY_LAYER_MODE
from ppt_system.export.ppt_text_style_rebuilder import rebuild_existing_ppt_text_styles


class PptTextStyleRebuilderTests(unittest.TestCase):
    def test_rebuild_changes_text_without_adding_images(self):
        with tempfile.TemporaryDirectory() as directory:
            job_dir = Path(directory)
            path = job_dir / EDITABLE_PPT_FILENAMES[OVERLAY_LAYER_MODE]
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            textbox.text_frame.paragraphs[0].add_run().text = "测试"
            presentation.save(path)

            rebuilt = rebuild_existing_ppt_text_styles(job_dir, 1, "page_text_style", {"style": {"font_name": "Microsoft YaHei", "color": "163A63"}})
            self.assertEqual(rebuilt, [path])
            result = Presentation(path)
            run = result.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
            self.assertEqual(run.font.name, "Microsoft YaHei")
            self.assertEqual(str(run.font.color.rgb), "163A63")


if __name__ == "__main__":
    unittest.main()
