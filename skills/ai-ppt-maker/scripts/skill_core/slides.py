from __future__ import annotations

import json
import os
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from asset_assembly import add_assets_to_slide
from ppt_system.export.export_artifact_policy import save_final_presentation_atomically
from ppt_system.export.reference_preview_export import export_reference_images_to_pptx
from ppt_system.export.text_style_runtime import apply_run_font_family, should_wrap_text
from ppt_system.export.ppt_calibration_renderer import render_pptx_first_slide_to_png

from skill_core.images import selected_pages


DELIVERY_MODES = {"image_only", "overlay", "separate_slides"}


def _color(value: Any, default: str = "FFFFFF") -> RGBColor:
    raw = str(value or default).strip().lstrip("#").upper()
    if len(raw) != 6 or any(char not in "0123456789ABCDEF" for char in raw):
        raise ValueError(f"颜色必须是六位十六进制值：{value}")
    return RGBColor.from_string(raw)


def _new_slide(presentation: Presentation, background_color: str):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _color(background_color)
    return slide


def _add_assets(slide, manifest_path: Path, slide_width: float, slide_height: float) -> int:
    if not manifest_path.is_file():
        raise FileNotFoundError(f"元素清单不存在：{manifest_path}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    for asset in manifest.get("assets", []):
        if not (manifest_path.parent / str(asset.get("file", ""))).is_file():
            raise FileNotFoundError(f"元素文件不存在：{asset.get('file')}")
    return add_assets_to_slide(slide, manifest_path, slide_width, slide_height)


def _add_text_box(slide, box: dict[str, Any], canvas_width: int, canvas_height: int, slide_width: float, slide_height: float) -> None:
    left = Inches(float(box["left"]) / canvas_width * slide_width)
    top = Inches(float(box["top"]) / canvas_height * slide_height)
    width = Inches(float(box["width"]) / canvas_width * slide_width)
    height = Inches(float(box["height"]) / canvas_height * slide_height)
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.fill.background()
    shape.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = Pt(0)
    frame.margin_right = Pt(0)
    frame.margin_top = Pt(0)
    frame.margin_bottom = Pt(0)
    frame.vertical_anchor = getattr(MSO_ANCHOR, str(box.get("valign", "TOP")).upper(), MSO_ANCHOR.TOP)
    frame.word_wrap = should_wrap_text(str(box["text"]), float(box["width"]), float(box["height"]), float(box["font_size"]))

    for line_index, line in enumerate(str(box["text"]).split("\n")):
        paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
        paragraph.alignment = getattr(PP_ALIGN, str(box.get("align", "LEFT")).upper(), PP_ALIGN.LEFT)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.0
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(float(box["font_size"]))
        run.font.bold = bool(box.get("bold", False))
        run.font.italic = bool(box.get("italic", False))
        run.font.color.rgb = _color(box.get("color"), "222222")
        apply_run_font_family(run, str(box.get("font_name", "Microsoft YaHei")))


def build_presentation(
    project: dict[str, Any],
    work_dir: Path,
    output_path: Path,
    *,
    mode: str = "separate_slides",
    page_numbers: list[int] | None = None,
) -> dict[str, Any]:
    if mode not in DELIVERY_MODES:
        raise ValueError("不支持的交付模式：" + mode)
    pages = selected_pages(project, page_numbers)
    canvas_width = int(project["canvas"]["width"])
    canvas_height = int(project["canvas"]["height"])
    slide_width = float(project.get("slide_width_inch", 13.333333))
    slide_height = slide_width * canvas_height / canvas_width

    if mode == "image_only":
        reference_pages = [
            {
                "page_no": int(page["page_no"]),
                "image": str((work_dir / f"page_{int(page['page_no']):02d}" / "reference.png").resolve()),
            }
            for page in pages
        ]
        result = export_reference_images_to_pptx(
            reference_pages,
            work_dir,
            output_path,
            image_width=canvas_width,
            image_height=canvas_height,
            slide_width_inch=slide_width,
        )
        return {"output_pptx": str(Path(result["pptx_path"]).resolve()), "logical_pages": len(pages), "slides": len(pages), "mode": mode}

    presentation = Presentation()
    presentation.slide_width = Inches(slide_width)
    presentation.slide_height = Inches(slide_height)
    for page in pages:
        page_no = int(page["page_no"])
        manifest_path = work_dir / f"page_{page_no:02d}" / "assets" / "assets.json"
        background = str(page.get("background_color") or project.get("background_color") or "#FFFFFF")
        if mode == "separate_slides":
            asset_slide = _new_slide(presentation, background)
            _add_assets(asset_slide, manifest_path, slide_width, slide_height)
            text_slide = _new_slide(presentation, background)
        else:
            text_slide = _new_slide(presentation, background)
            _add_assets(text_slide, manifest_path, slide_width, slide_height)
        for box in page.get("text_boxes", []):
            _add_text_box(text_slide, box, canvas_width, canvas_height, slide_width, slide_height)

    saved = save_final_presentation_atomically(presentation, output_path)
    return {
        "output_pptx": str(saved.resolve()),
        "logical_pages": len(pages),
        "slides": len(presentation.slides),
        "mode": mode,
    }


def render_page_previews(
    project: dict[str, Any],
    work_dir: Path,
    output_dir: Path,
    *,
    page_numbers: list[int] | None = None,
) -> list[dict[str, Any]]:
    if os.name != "nt":
        raise RuntimeError("PowerPoint 真实渲染需要 Windows 与本机 PowerPoint。")
    width = int(project["canvas"]["width"])
    height = int(project["canvas"]["height"])
    previews: list[dict[str, Any]] = []
    for page in selected_pages(project, page_numbers):
        page_no = int(page["page_no"])
        preview_pptx = output_dir / f"page_{page_no:02d}.pptx"
        preview_png = output_dir / f"page_{page_no:02d}.png"
        build_presentation(project, work_dir, preview_pptx, mode="overlay", page_numbers=[page_no])
        rendered = render_pptx_first_slide_to_png(
            preview_pptx,
            preview_png,
            image_width=width,
            image_height=height,
        )
        if rendered is None:
            raise RuntimeError("本机 PowerPoint 真实渲染失败；请确认 Windows PowerPoint 可用。预览 PPTX 已保留。")
        previews.append({"page_no": page_no, "preview_pptx": str(preview_pptx.resolve()), "preview_png": str(rendered.resolve())})
    return previews
