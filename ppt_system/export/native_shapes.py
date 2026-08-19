from __future__ import annotations

from collections.abc import Callable, Mapping
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt


ShapeScaler = Callable[[float], Any]

_DEFAULT_FILL = "E8EEF7"
_DEFAULT_LINE = "8EA3BF"
_DECORATIVE_FILL = "D7E3F2"
_ACCENT_FILL = "DCEBFA"

_KIND_SHAPES = {
    "circle": MSO_AUTO_SHAPE_TYPE.OVAL,
    "circle-center": MSO_AUTO_SHAPE_TYPE.OVAL,
    "hub": MSO_AUTO_SHAPE_TYPE.OVAL,
    "hub-node": MSO_AUTO_SHAPE_TYPE.OVAL,
    "line-node": MSO_AUTO_SHAPE_TYPE.OVAL,
    "dot": MSO_AUTO_SHAPE_TYPE.OVAL,
    "map-pin": MSO_AUTO_SHAPE_TYPE.OVAL,
    "ring": MSO_AUTO_SHAPE_TYPE.DONUT,
    "pie": MSO_AUTO_SHAPE_TYPE.PIE,
    "triangle": MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE,
    "pyramid": MSO_AUTO_SHAPE_TYPE.TRAPEZOID,
    "funnel": MSO_AUTO_SHAPE_TYPE.TRAPEZOID,
    "chevron": MSO_AUTO_SHAPE_TYPE.CHEVRON,
    "arrow-node": MSO_AUTO_SHAPE_TYPE.CHEVRON,
    "arrow-left": MSO_AUTO_SHAPE_TYPE.LEFT_ARROW,
    "arrow-right": MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
    "flag": MSO_AUTO_SHAPE_TYPE.WAVE,
    "check": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
}

_LINE_KINDS = {"line-h", "line-v", "axis-x", "axis-y", "trend-line", "water", "bone-main"}


def add_native_shapes(
    slide: Any,
    page: Mapping[str, Any],
    scale_x: ShapeScaler,
    scale_y: ShapeScaler,
    scale_width: ShapeScaler,
    scale_height: ShapeScaler,
) -> list[Any]:
    """把页面蓝图转换成可编辑的 PowerPoint 原生形状。"""
    created: list[Any] = []
    blueprint = page.get("native_blueprint", [])
    if not isinstance(blueprint, list):
        return created

    for item in blueprint:
        if not isinstance(item, Mapping):
            continue
        try:
            left = float(item.get("left", 0))
            top = float(item.get("top", 0))
            width = float(item.get("width", 0))
            height = float(item.get("height", 0))
        except (TypeError, ValueError):
            continue
        if width <= 0 or height <= 0:
            continue

        kind = str(item.get("kind") or "rect").strip().lower()
        shape_type = _KIND_SHAPES.get(kind, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
        shape = slide.shapes.add_shape(
            shape_type,
            scale_x(left),
            scale_y(top),
            scale_width(width),
            scale_height(height),
        )
        _style_shape(shape, item, kind)
        created.append(shape)
    return created


def _style_shape(shape: Any, item: Mapping[str, Any], kind: str) -> None:
    decorative = bool(item.get("decorative"))
    fill_color = _normalize_color(
        item.get("fill"),
        _DECORATIVE_FILL if decorative else (_ACCENT_FILL if kind in {"hub", "circle-center", "big-number"} else _DEFAULT_FILL),
    )
    line_color = _normalize_color(item.get("line"), _DEFAULT_LINE)

    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill_color)
    shape.line.color.rgb = RGBColor.from_string(line_color)
    shape.line.width = Pt(1)
    if kind in _LINE_KINDS:
        shape.line.fill.background()

    label = str(item.get("label") or "").strip()
    if decorative or not label or not getattr(shape, "has_text_frame", False):
        return
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.margin_left = Pt(4)
    text_frame.margin_right = Pt(4)
    text_frame.margin_top = Pt(2)
    text_frame.margin_bottom = Pt(2)
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = label
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string("27415F")


def _normalize_color(value: Any, default: str) -> str:
    normalized = str(value or "").strip().lstrip("#").upper()
    if len(normalized) != 6 or any(character not in "0123456789ABCDEF" for character in normalized):
        return default
    return normalized
