"""结构化、可编辑的 PowerPoint 图表数据与导出辅助。"""
from __future__ import annotations

from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

CHART_TYPES = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
               "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE,
               "doughnut": XL_CHART_TYPE.DOUGHNUT, "area": XL_CHART_TYPE.AREA}

def normalize_chart_data(value: Any) -> dict[str, Any] | None:
    if isinstance(value, str):
        rows = []
        for line in value.splitlines():
            line = line.strip()
            if not line or line.startswith("#"): continue
            parts = [item.strip() for item in line.replace("：", ":").replace("，", ",").split(",", 1)]
            if len(parts) == 1: parts = [item.strip() for item in line.split(":", 1)]
            if len(parts) != 2: raise ValueError("图表数据每行格式应为：分类,数值")
            rows.append((parts[0], parts[1]))
        value = {"type": "bar", "categories": [row[0] for row in rows], "series": [{"name": "数值", "values": [row[1] for row in rows]}]}
    if not isinstance(value, dict):
        return None
    chart_type = str(value.get("type", "bar")).strip().lower()
    if chart_type not in CHART_TYPES:
        raise ValueError(f"不支持的图表类型：{chart_type}")
    categories = [str(item).strip() for item in value.get("categories", []) if str(item).strip()]
    series = []
    for raw in value.get("series", []):
        if not isinstance(raw, dict):
            continue
        values = []
        for item in raw.get("values", []):
            try: values.append(float(item))
            except (TypeError, ValueError) as exc: raise ValueError("图表数据必须全部为数字") from exc
        if len(values) != len(categories):
            raise ValueError("图表分类数量必须与数据数量一致")
        series.append({"name": str(raw.get("name", "系列")).strip() or "系列", "values": values})
    if not categories or not series: return None
    return {"type": chart_type, "title": str(value.get("title", "")).strip(),
            "categories": categories, "series": series,
            "show_legend": bool(value.get("show_legend", len(series) > 1)),
            "show_data_labels": bool(value.get("show_data_labels", False)),
            "left": max(0, int(value.get("left", 1000))), "top": max(0, int(value.get("top", 260))),
            "width": max(1, int(value.get("width", 900))), "height": max(1, int(value.get("height", 650)))}

def add_editable_chart(slide, chart_data: dict[str, Any], x, y, w, h):
    normalized = normalize_chart_data(chart_data)
    if not normalized: return None
    data = CategoryChartData(); data.categories = normalized["categories"]
    for series in normalized["series"]: data.add_series(series["name"], series["values"])
    chart = slide.shapes.add_chart(CHART_TYPES[normalized["type"]], x, y, w, h, data).chart
    chart.has_legend = normalized["show_legend"]
    if normalized["title"]:
        chart.has_title = True; chart.chart_title.text_frame.text = normalized["title"]
    if normalized["show_data_labels"]: chart.plots[0].has_data_labels = True
    return chart
