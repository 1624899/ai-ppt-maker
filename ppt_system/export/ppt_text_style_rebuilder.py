from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from ppt_system.export.delivery_options import EDITABLE_PPT_FILENAMES
from ppt_system.export.export_layer_mode import SEPARATE_LAYER_MODE


def rebuild_existing_ppt_text_styles(job_dir: Path, page_no: int, operation_type: str, payload: dict[str, Any]) -> list[Path]:
    """直接修改已有 PPTX 文字对象，不调用模型或重新生成图片。"""
    rebuilt: list[Path] = []
    for layer_mode, filename in EDITABLE_PPT_FILENAMES.items():
        path = job_dir / filename
        if not path.exists():
            continue
        presentation = Presentation(str(path))
        slide_indexes = _slide_indexes(page_no, layer_mode, len(presentation.slides))
        for slide_index in slide_indexes:
            _apply_to_slide(presentation.slides[slide_index], operation_type, payload)
        temporary = path.with_suffix(path.suffix + ".updating")
        presentation.save(str(temporary))
        os.replace(temporary, path)
        rebuilt.append(path)
    return rebuilt


def _slide_indexes(page_no: int, layer_mode: str, slide_count: int) -> list[int]:
    logical_index = max(0, int(page_no) - 1)
    candidates = [logical_index * 2 + 1] if layer_mode == SEPARATE_LAYER_MODE else [logical_index]
    return [index for index in candidates if index < slide_count]


def _apply_to_slide(slide, operation_type: str, payload: dict[str, Any]) -> None:
    style = payload.get("style") if isinstance(payload.get("style"), dict) else {}
    align = getattr(PP_ALIGN, str(payload.get("align", "LEFT")).upper(), PP_ALIGN.LEFT)
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            if operation_type == "page_align":
                paragraph.alignment = align
            for run in paragraph.runs:
                if "font_name" in style:
                    run.font.name = str(style["font_name"])
                if "font_size" in style:
                    from pptx.util import Pt
                    run.font.size = Pt(float(style["font_size"]))
                if "color" in style:
                    color = str(style["color"]).replace("#", "").upper()
                    if len(color) == 6:
                        run.font.color.rgb = RGBColor.from_string(color)
                if "bold" in style:
                    run.font.bold = bool(style["bold"])
                if "italic" in style:
                    run.font.italic = bool(style["italic"])
